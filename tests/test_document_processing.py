"""
Unit & Integration Tests for Document Processing Modules
"""

import os
import unittest
import tempfile
import pymupdf
from docx import Document

from backend.document_processing.pdf_extractor import extract_pdf_text
from backend.document_processing.docx_extractor import extract_docx_text
from backend.document_processing.ocr import extract_image_text, is_ocr_available


class TestDocumentProcessing(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_extract_pdf_valid(self):
        """Test standard PDF text extraction."""
        pdf_path = os.path.join(self.temp_dir.name, "sample.pdf")
        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((50, 50), "This is a valid legal contract for testing.")
        doc.save(pdf_path)
        doc.close()

        text = extract_pdf_text(pdf_path)
        self.assertIn("This is a valid legal contract", text)

    def test_extract_pdf_missing_file(self):
        """Test handling of non-existent PDF file."""
        with self.assertRaises(FileNotFoundError):
            extract_pdf_text("non_existent_file_12345.pdf")

    def test_extract_docx_valid(self):
        """Test DOCX text extraction with paragraphs and tables."""
        docx_path = os.path.join(self.temp_dir.name, "sample.docx")
        doc = Document()
        doc.add_paragraph("Employment Agreement between ABC Corp and John Doe.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Monthly Salary"
        table.cell(0, 1).text = "INR 50,000"
        doc.save(docx_path)

        text = extract_docx_text(docx_path)
        self.assertIn("Employment Agreement", text)
        self.assertIn("Monthly Salary", text)
        self.assertIn("INR 50,000", text)

    def test_extract_docx_missing_file(self):
        """Test handling of non-existent DOCX file."""
        with self.assertRaises(FileNotFoundError):
            extract_docx_text("non_existent_file_12345.docx")

    def test_ocr_availability_check(self):
        """Test that is_ocr_available returns boolean without throwing uncaught exceptions."""
        avail = is_ocr_available()
        self.assertIsInstance(avail, bool)


if __name__ == "__main__":
    unittest.main()
