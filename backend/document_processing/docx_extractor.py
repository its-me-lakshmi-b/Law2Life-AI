import os
from docx import Document


def extract_docx_text(file_path: str) -> str:
    """
    Extract text from a DOCX document including paragraphs and tables.

    Args:
        file_path (str): Path to the DOCX file.

    Returns:
        str: Extracted text from the DOCX document.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    text_parts = []

    try:
        document = Document(file_path)

        for paragraph in document.paragraphs:
            if paragraph.text and paragraph.text.strip():
                text_parts.append(paragraph.text.strip())

        # Also extract text from any tables inside DOCX
        for table in document.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Remove consecutive duplicates in merged cells
                    unique_cells = []
                    for c in row_cells:
                        if not unique_cells or c != unique_cells[-1]:
                            unique_cells.append(c)
                    text_parts.append(" | ".join(unique_cells))

        return "\n\n".join(text_parts).strip()

    except Exception as e:
        raise Exception(f"Error extracting DOCX text: {str(e)}")